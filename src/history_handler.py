"""协议测试历史的 CSV / Excel / JSON 交换，兼容旧版导出列。"""
from __future__ import annotations

import csv
import io
import re
import json
from datetime import datetime
from pathlib import Path

from src.database import ProtocolTestSession
from src.protocol_history import load_protocol_history_params


HEADERS = ['测试时间', '协议', '目标IP', '端口', '结果', '发送报文', '响应报文',
           '错误信息', '测试参数', '发送原始字节', '响应原始字节']
REQUIRED = set(HEADERS[:5])
PROTOCOLS = {'TCP': 'tcp_client', 'WS': 'ws_client', 'HTTP': 'http_client',
             'TCP_CLIENT': 'tcp_client', 'WS_CLIENT': 'ws_client', 'HTTP_CLIENT': 'http_client'}


def _text(value):
    return '' if value is None else str(value)


def _hex_bytes(value):
    value = _text(value).strip()
    if not value:
        return None
    return bytes.fromhex(value[2:] if value.startswith('0x') else value)


def _parse_record(row: dict) -> ProtocolTestSession:
    if not REQUIRED <= row.keys():
        raise ValueError('缺少列：' + '、'.join(sorted(REQUIRED - row.keys())))
    proto = PROTOCOLS.get(_text(row['协议']).strip().upper())
    if not proto:
        raise ValueError('协议必须为 TCP、WS 或 HTTP')
    result = _text(row['结果']).strip().upper()
    if result not in ('OK', 'FAIL'):
        raise ValueError('结果必须为 OK 或 FAIL')
    value = row['测试时间']
    timestamp = value.isoformat(sep=' ') if isinstance(value, datetime) else _text(value).strip()
    datetime.fromisoformat(timestamp)
    port_value = _text(row['端口']).strip() or '0'
    port = int(port_value)
    if not 0 <= port <= 65535:
        raise ValueError('端口应在 0 至 65535 之间')
    params = row.get('测试参数')
    if isinstance(params, str) and params.strip():
        params = json.loads(params)
    if params in (None, ''):
        params = {}
    if not isinstance(params, dict):
        raise ValueError('测试参数必须是 JSON 对象')
    serialized = json.dumps(params, ensure_ascii=False)
    serialized = json.dumps(load_protocol_history_params(serialized, proto), ensure_ascii=False)
    return ProtocolTestSession(
        id=0, protocol_type=proto, started_at=timestamp,
        target_ip=_text(row['目标IP']), target_port=port, success=result == 'OK',
        request=_text(row.get('发送报文')), response=_text(row.get('响应报文')),
        error_msg=_text(row.get('错误信息')), test_params=serialized,
        request_raw=_hex_bytes(row.get('发送原始字节')),
        response_raw=_hex_bytes(row.get('响应原始字节')))


def _read_csv_text(text: str) -> list[dict]:
    """CSV 读取：Python csv 拒绝 NUL，按反斜转转序列在内存转后恢复。"""
    escaped = '\x00' in text
    if escaped:
        text = text.replace('\\', '\\\\').replace('\x00', '\\0')
    rows = list(csv.DictReader(io.StringIO(text)))
    if not escaped:
        return rows
    return [{key: re.sub(r'\\([\\0])',
                    lambda m: '\x00' if m.group(1) == '0' else '\\', value)
                    if isinstance(value, str) else value
             for key, value in row.items()} for row in rows]


def read_history(filepath) -> list[ProtocolTestSession]:
    path = Path(filepath)
    if path.suffix.lower() == '.json':
        document = json.loads(path.read_text(encoding='utf-8-sig'))
        if not isinstance(document, dict) or document.get('format') != 'TestTool.ProtocolHistory' \
                or document.get('version') != 1 or not isinstance(document.get('records'), list):
            raise ValueError('不是 TestTool 测试历史 JSON 文件')
        records = document['records']
    elif path.suffix.lower() == '.csv':
        csv.field_size_limit(2 ** 31 - 1)
        with path.open(encoding='utf-8-sig', newline='') as source:
            records = _read_csv_text(source.read())
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [_text(cell).strip() for cell in next(rows, ())]
            records = [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()
    else:
        raise ValueError('仅支持 .csv、.xlsx 和 .json 测试历史文件')
    sessions = []
    for index, row in enumerate(records, 2):
        if not isinstance(row, dict):
            raise ValueError(f'第 {index} 行：记录必须为对象')
        if all(value in (None, '') for value in row.values()):
            continue
        try:
            sessions.append(_parse_record(row))
        except (ValueError, TypeError, OverflowError) as exc:
            raise ValueError(f'第 {index} 行：{exc}') from exc
    if not sessions:
        raise ValueError('文件中没有测试历史记录')
    return sessions


def write_history(filepath, sessions: list[ProtocolTestSession]):
    path = Path(filepath)
    rows = [[
        s.started_at, {'tcp_client': 'TCP', 'ws_client': 'WS', 'http_client': 'HTTP'}[s.protocol_type],
        s.target_ip, s.target_port, 'OK' if s.success else 'FAIL', s.request or '', s.response or '',
        s.error_msg or '', json.dumps(load_protocol_history_params(s.test_params, s.protocol_type), ensure_ascii=False),
        None if s.request_raw is None else '0x' + s.request_raw.hex(),
        None if s.response_raw is None else '0x' + s.response_raw.hex(),
    ] for s in sessions]
    if path.suffix.lower() == '.json':
        document = dict(format='TestTool.ProtocolHistory', version=1,
                        records=[dict(zip(HEADERS, row)) for row in rows])
        path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding='utf-8')
    elif path.suffix.lower() == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as destination:
            writer = csv.writer(destination)
            writer.writerow(HEADERS)
            writer.writerows(rows)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import Workbook
        # Excel 会截断过长文本，或无法保存控制字符；完整报文应选择 CSV/JSON。
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        if any(isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value))
               for row in rows for value in row):
            raise ValueError('报文超出 Excel 文本限制，请使用 CSV 或 JSON 导出以保留完整内容')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '测试历史'
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        workbook.save(path)
        workbook.close()
    else:
        raise ValueError('仅支持 .csv、.xlsx 和 .json 测试历史文件')
